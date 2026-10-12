# flask_app.py

# ==========================================
# 第 1 區：基礎環境與系統設定
# ==========================================

from flask import Flask, request, render_template, abort, url_for, session, redirect, make_response, jsonify
import csv
import os
import requests
import time
from datetime import datetime, timedelta, timezone
from thefuzz import fuzz
from authlib.integrations.flask_client import OAuth
from werkzeug.middleware.proxy_fix import ProxyFix
from dotenv import load_dotenv
from flask_sqlalchemy import SQLAlchemy
import calendar
from functools import wraps
from linebot import LineBotApi, WebhookHandler
from linebot.exceptions import InvalidSignatureError
from linebot.models import MessageEvent, TextMessage, TextSendMessage
from flask_wtf.csrf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from werkzeug.utils import secure_filename

# 校慶紀念服
import openpyxl
from openpyxl.worksheet.pagebreak import Break # 分頁套件
from openpyxl.styles import PatternFill, Alignment, Border, Side, Font
from itertools import groupby
from collections import defaultdict
from flask import send_file


try:
    from school_news import get_news_html
except ImportError:
    def get_news_html(): return "爬蟲模組載入失敗"

BASE_DIR = './mysite'
DATA_DIR = os.path.join(BASE_DIR, 'data')
CSV_PATH_TEACHERS = os.path.join(DATA_DIR, 'teachers.csv')
CSV_PATH_WEDNESDAY = os.path.join(DATA_DIR, 'wednesday.csv')

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1) # 使用者的連線不會直接碰到你的 Flask 程式

load_dotenv(os.path.join(BASE_DIR, '.env')) # 密碼存放在.env中
app.secret_key = os.getenv('FLASK_SECRET_KEY') or 'hshs_secret_12345' # 要有key才能使用session

app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///app.db' # 指定資料庫型態
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False # 關閉資料庫追蹤功能
# 第二個資料庫
app.config['SQLALCHEMY_BINDS'] = {
    'order_db': 'sqlite:///order.db'
}

csrf = CSRFProtect(app) # 啟動全站的 CSRF 保護

limiter = Limiter(
    get_remote_address,  # 使用者的 IP 地址
    app=app,
    default_limits=["200 per day", "50 per hour"], # 預設：全站每個 IP 一天最多 200 次請求，一小時 50 次
    storage_uri="memory://" # 將記錄存在記憶體中
)

# ==========================================
# 第 2 區：資料庫興建
# ==========================================

db = SQLAlchemy(app)

class DishComment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    shop = db.Column(db.String(50), nullable=False)
    dish_name = db.Column(db.String(100), nullable=False)
    student_class = db.Column(db.String(20), nullable=False)
    comment = db.Column(db.Text, nullable=False)

class MainComment(db.Model):
    id = db.Column(db.Integer,primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    to = db.Column(db.String(20), nullable=False)
    comment = db.Column(db.Text, nullable=False)
    check = db.Column(db.Integer, nullable=False) # 0為沒處理，1為已處理

class BandaiSignup(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    seat = db.Column(db.String(20), nullable=False)
    name = db.Column(db.String(20), nullable=False)
    check = db.Column(db.Integer, nullable=False)

class LineMessage(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    name = db.Column(db.String(100), nullable=False)
    message = db.Column(db.Text, nullable=False)
    check = db.Column(db.Integer, nullable=False)

class Counter(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(50), unique=True, nullable=False)
    total_count = db.Column(db.Integer, default=0)

class MisTake(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    area = db.Column(db.String(50),nullable=False)
    message = db.Column(db.Text,nullable=False)
    check = db.Column(db.Integer, default=0)

class InterviewWish(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    choose = db.Column(db.String(50), default='無')
    interviewee = db.Column(db.String(50), default='無')
    reason = db.Column(db.String(50), default='無')
    topic = db.Column(db.Text,nullable=False)

class ChooseSong(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    date = db.Column(db.String(20), nullable=False)
    song_name = db.Column(db.String(50),nullable=False)
    artist = db.Column(db.String(50), default='無')
    yt_link = db.Column(db.Text, default='無')


# 訂購平台 人
class OrderUser(db.Model):
    __bind_key__ = 'order_db'
    __tablename__ = 'order_user'
    id = db.Column(db.Integer, primary_key=True)
    is_student = db.Column(db.Boolean, nullable=False)
    login_id = db.Column(db.String(50), unique=True, nullable=False)
    class_name = db.Column(db.String(50), nullable=False)
    seat_number = db.Column(db.String(50), nullable=False)
    student_name = db.Column(db.String(50), nullable=False)

    orders = db.relationship('OrderMain', backref='user', lazy=True)

class OrderMain(db.Model):
    __bind_key__ = 'order_db'
    __tablename__ = 'order_main'
    order_id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('order_user.id'), nullable=False)
    status = db.Column(db.String(50), default='購物車', nullable=False)
    items = db.relationship('OrderItem', backref='order', lazy=True)


class OrderItem(db.Model):
    __bind_key__ = 'order_db'
    __tablename__ = 'order_item'
    item_id = db.Column(db.Integer, primary_key=True)
    order_id = db.Column(db.Integer, db.ForeignKey('order_main.order_id'), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey('order_product.id'), nullable=False)
    item_size = db.Column(db.String(50), nullable=False)
    quantity = db.Column(db.Integer, default=1, nullable=False)
    product = db.relationship('OrderProduct')

# 訂購平台 物
class OrderProduct(db.Model):
    __bind_key__ = 'order_db'
    __tablename__ = 'order_product'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    price = db.Column(db.Integer, nullable=False)
    images = db.Column(db.Text, nullable=False)
    sizes = db.Column(db.Text, nullable=False)

class OrderSetting(db.Model):
    __bind_key__ = 'order_db'
    __tablename__ = 'order_setting'
    id = db.Column(db.Integer, primary_key=True)
    start_time = db.Column(db.String(20), nullable=True) # 格式為 YYYY-MM-DDTHH:MM
    end_time = db.Column(db.String(20), nullable=True)
    marquee = db.Column(db.Text, nullable=True)


with app.app_context():
    db.create_all()
    if not OrderSetting.query.first():
            db.session.add(OrderSetting())
            db.session.commit()
    initial_counts = {
        'dish_comment': 0,
        'main_comment': 0,
        'bandai_signup': 0,
        'line_message': 0,
        'interview_wish':0
    }
    for name, start_value in initial_counts.items():
        if not Counter.query.filter_by(name=name).first():
            new_counter = Counter(name=name, total_count=start_value)
            db.session.add(new_counter)
    db.session.commit()


def add_count(counter_name, increment=False):
    counter = Counter.query.filter_by(name=counter_name).first()
    if not counter:
        counter = Counter(name=counter_name, total_count=0)
        db.session.add(counter)
        db.session.commit()
    if increment:
        try:
            counter.total_count += 1
            db.session.commit()
        except:
            db.session.rollback()
    return counter.total_count

# ==========================================
# 第 3 區：LINE Bot 官方帳號系統
# ==========================================

line_token = os.getenv('LINE_CHANNEL_ACCESS_TOKEN')
line_secret = os.getenv('LINE_CHANNEL_SECRET')

with app.app_context():
    try:
        line_bot_api = LineBotApi((line_token or 'missing_token').strip())
        webhook_handler = WebhookHandler((line_secret or 'missing_secret').strip())
    except Exception as e:
        line_bot_api = None
        webhook_handler = None
        db.session.add(MisTake(
            area = "LINE Bot 初始化",
            message = f"LINE Bot 實例化失敗，請檢查 Token 格式。錯誤：{e}"
        ))
        db.session.commit()

    if not (line_token and line_secret):
        db.session.add(MisTake(
            area = "LINE 金鑰讀取",
            message = "系統無法從 .env 讀取到 LINE_CHANNEL_ACCESS_TOKEN 或 LINE_CHANNEL_SECRET。"
        ))
        db.session.commit()

def get_weather_forecast():
    try:
        data = requests.get("https://api.open-meteo.com/v1/forecast?latitude=25.01&longitude=121.46&daily=temperature_2m_max,temperature_2m_min,precipitation_probability_max&timezone=Asia%2FTaipei", timeout=5).json()
        return f"【海山高中今日氣象】\n📍 板橋區預測：\n🌡️ 最低溫: {data['daily']['temperature_2m_min'][0]}°C\n🌡️ 最高溫: {data['daily']['temperature_2m_max'][0]}°C\n☔ 下雨機率: {data['daily']['precipitation_probability_max'][0]}%\n\n祝你有個美好的一天！"
    except Exception as e:
        db.session.add(MisTake(
                area = "今日天氣",
                message = f"有學生暫時無法取得即時天氣資料，電腦給的錯誤為：{e}"
            )
        )
        db.session.commit()
        return "暫時無法取得即時天氣資料，大多數的原因為：同時多人點擊，我們的LINE沒有付費，造成使用不便不好意思，請稍後再試!也可以請發送 【錯誤】 提醒我們!感謝"


def get_news(choose):
    API_KEY = os.getenv('NEWS_API_KEY')
    if not API_KEY:
        db.session.add(MisTake(
                area = "今日新聞",
                message = "系統未設定新聞 API 金鑰，API 金鑰存放在.env檔案，沒有這個可能造成【今日天氣】【今日新聞】【校網20天內重大公告】都失效。"
            )
        )
        db.session.commit()
        return "錯誤!(系統未設定新聞 API 金鑰)，請發送 【錯誤】 提醒我們!感謝"


    def fetch_news(category):
        if category == "台灣":
            search_query = "台灣 AND (教育 OR 政治 OR 社會 OR 科技 OR 經濟 OR 文化)"
        else:
            search_query = "國際 OR 全球"

        url = f"https://newsapi.org/v2/everything?q={search_query}&language=zh&sortBy=publishedAt&pageSize=5&apiKey={API_KEY}"

        try:
            headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
            response = requests.get(url, headers=headers, timeout=5)
            data = response.json()

            if data.get('status') == 'error':
                db.session.add(MisTake(
                    area = "今日新聞",
                    message = f"NewsAPI 發生錯誤，電腦給的錯誤為：({category}): {data.get('message')}"
                ))
                db.session.commit()
                return []

            return data.get('articles', [])

        except Exception as e:
            db.session.add(MisTake(
                area = "今日新聞",
                message = f"網路連線或解析失敗，電腦給的錯誤為：({category}): {e}"
            ))
            db.session.commit()
            return []

    tw_news = []
    int_news = []

    if choose in ["taiwan", "all"]:
        tw_news = fetch_news("台灣")

    if choose in ["world", "all"]:
        int_news = fetch_news("國際")

    if not tw_news and not int_news:
        return "錯誤!今日暫時無法取得新聞，請稍後再試！請發送 【錯誤】提醒我們!謝謝!"

    final_report = "📰 【今日重點新聞】\n\n"

    if int_news:
        final_report += "🌍 國際焦點：\n"
        for idx, art in enumerate(int_news, 1):
            title = art.get('title', '無標題')
            desc = art.get('description') or '無提供摘要'
            desc_short = desc[:200].replace('\n', '') + ("..." if len(desc) > 200 else "")
            final_report += f"{idx}. {title}\n📝 {desc_short}\n\n"

    if tw_news:
        final_report += "🇹🇼 台灣時事：\n"
        for idx, art in enumerate(tw_news, 1):
            title = art.get('title', '無標題')
            desc = art.get('description') or '無提供摘要'
            desc_short = desc[:200].replace('\n', '') + ("..." if len(desc) > 200 else "")
            final_report += f"{idx}. {title}\n📝 {desc_short}\n\n"

    return final_report.strip()


def get_wednesday(grade):
    if not os.path.exists(CSV_PATH_WEDNESDAY):
        db.session.add(MisTake(
            area = "週三行程查詢",
            message = f"找不到週三行程檔案：wednesday.csv，代表指定位置沒有命名正確的檔案"
        ))
        db.session.commit()
        return "錯誤!(找不到行程檔案)，請發送 【錯誤】提醒我們。"
    today = datetime.now()
    current_weekday = today.weekday()

    if current_weekday <= 2:
        days_to_target = 2 - current_weekday
    else:
        days_to_target = (2 - current_weekday) + 7

    target_wednesday = today + timedelta(days=days_to_target)
    target_date_str = f"{target_wednesday.month}/{target_wednesday.day}"

    try:
        with open(CSV_PATH_WEDNESDAY, mode='r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                if row.get('日期', '').strip() == target_date_str:
                    week_num = row.get('週次', '')
                    p5 = row.get('5', '')
                    if grade == "高一":
                        p67 = row.get('高一', '')
                    else:
                        p67 = row.get('高二', '')
                    return f"📅 查詢日期：{target_date_str}\n【第{week_num}週】\n第5節：{p5}\n第6、7節：{p67}"
            db.session.add(MisTake(
                area = "週三行程查詢",
                message = f" 【不一定錯誤】 找不到{target_date_str} 的排程，有可能真的沒有這周也有可能是檔案打錯日期"
            ))
            db.session.commit()
            return f"目前找不到 {target_date_str} 的排程資料喔！"

    except Exception as e:
        db.session.add(MisTake(
            area = "週三行程查詢",
            message = f"讀取發生錯誤：{str(e)}"
        ))
        db.session.commit()
        return "讀取行程檔案時發生錯誤，請稍後再試！"

def get_menu(input_file):
    now = datetime.now(timezone(timedelta(hours=8)))
    today = f"{now.month}/{now.day}"
    file_name = f"{now.year}年{now.month}月海山高中菜單_{input_file}.csv"
    full_path = os.path.join(BASE_DIR, 'menus', file_name)
    if not os.path.exists(full_path):
        error_msg = f"找不到菜單檔案：{file_name}"
        try:
            db.session.add(MisTake(
                area = "今日菜色",
                message = error_msg
            ))
            db.session.commit()
        except Exception as db_e:
            db.session.rollback()
            error_msg += f"\n(資料庫記錄失敗: {str(db_e)})"
        return f"目前沒有提供午餐"
        # return f"錯誤！詳細錯誤：{error_msg}，請發送 【錯誤】提醒我們。"
    try:
        with open(full_path, mode='r', encoding='utf-8-sig') as f:
            for row in csv.DictReader(f):
                if row.get('日期', '').strip() == today:
                    return f"今天的菜色是：\n{row['品項']}"
        try:
            db.session.add(MisTake(
                area = "今日菜色",
                message = f"找不到{today}菜單，可能是放假或格式錯誤"
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()

        return "目前沒有提供午餐"
    except Exception as e:
        error_msg = f"讀取錯誤：{str(e)}"
        try:
            db.session.add(MisTake(
                area = "今日菜色",
                message = error_msg
            ))
            db.session.commit()
        except Exception as db_e:
            db.session.rollback()
            error_msg += f"\n(資料庫記錄失敗: {str(db_e)})"
        return f"發生系統錯誤！詳細原因：\n{error_msg}"

@app.route("/line/webhook", methods=['POST'])
@csrf.exempt
def line_webhook_callback():
    try:
        webhook_handler.handle(request.get_data(as_text=True), request.headers.get('X-Line-Signature', ''))
    except InvalidSignatureError:
        abort(400)
    return 'OK'

@webhook_handler.add(MessageEvent, message=TextMessage)
def handle_line_text_message(event):
    current_datetime = datetime.now(timezone(timedelta(hours=8)))
    formatted_time = current_datetime.strftime("%Y-%m-%d %H:%M:%S")
    received_text = event.message.text

    reply_text = ""

    if received_text == "今日天氣":
        reply_text = get_weather_forecast()
    elif received_text in ["素食", "統鮮", "上將", "宏遠"]:
        reply_text = get_menu(received_text)
    elif received_text == "國際台灣我都要":
        reply_text = get_news("all")
    elif received_text == "國際新聞":
        reply_text = get_news("world")
    elif received_text == "台灣新聞":
        reply_text = get_news("taiwan")
    elif received_text == "高一":
        reply_text = get_wednesday("高一")
    elif received_text == "高二":
        reply_text = get_wednesday("高二")
    elif received_text != f"{current_datetime.month}月菜色" and received_text != "今日新聞" and received_text != "行事曆" and received_text != "週三行程" and received_text != "今日菜色" and received_text != "我不知道營養午餐廠商" and received_text != "課表" and received_text != "本週三行程&今日新聞&今日天氣":
        user_id = event.source.user_id
        user_display_name = user_id
        try:
            user_display_name = line_bot_api.get_profile(user_id).display_name
        except:
            pass
        try:
            db.session.add(LineMessage(
                date=formatted_time,
                name=user_display_name,
                message=received_text,
                check=0
            ))
            db.session.commit()
            add_count('line_message', increment=True)
            reply_text = f"您好！班聯會已收到訊息，將盡快回復。\n也希望您能幫我們多加宣傳！"
        except Exception as e:
            db.session.rollback()
            db.session.add(MisTake(
                area = "資料庫LineMessage存檔",
                message = f"代表有學生發LINE訊息，但我們沒在後台網頁寫入。電腦給的錯誤為：資料庫錯誤: {e}"
            ))
            db.session.commit()
            reply_text = "系統發生一點小錯誤，儲存失敗，請再發一次。"

    if reply_text:
        line_bot_api.reply_message(event.reply_token, TextSendMessage(text=reply_text))

# ==========================================
# 第 4 區：學生服務中心(純路由)
# ==========================================

@app.route('/student/main')
def student_main():
    return render_template('student/main.html')

@app.route('/student/banlian_zzjs')
def student_banlian_zzjs():
    return render_template('student/banlian_zzjs.html')

@app.route('/student/date')
def student_date():
    return render_template('student/date.html')

@app.route('/student/main_comment')
def student_main_comment():
    return render_template('student/main_comment.html')

@app.route('/student/bandai')
def student_bandai():
    return render_template('student/bandai.html')

@app.route('/student/class_schedule')
def student_class_schedule():
    return render_template('student/class_schedule.html')

@app.route('/student/teacher_search')
def student_teacher_search():
    return render_template('student/teacher_search.html')

@app.route('/student/banlian_event.html')
def student_banlian_event():
    return render_template('student/banlian_event.html')

@app.route('/school_new')
def school_new_page():
    try: return get_news_html()
    except Exception as e:
        db.session.add(MisTake(
            area = "校網20天內重大公告",
            message = f"爬蟲執行出錯，電腦給的錯誤為：{str(e)}"
        ))
        db.session.commit()
    return "<h1>抱歉，校網公告讀取失敗</h1><p>已自動回報給系統管理員。</p>"

@app.route('/student/score_calculation')
def student_score_calculation():
    return render_template('student/score_calculation.html')

@app.route('/student/OOXXgame')
def student_OOXXgame():
    return render_template('student/OOXXgame.html')

@app.route('/student/pro_OOXXgame')
def student_pro_OOXXgame():
    return render_template('student/pro_OOXXgame.html')

@app.route('/student/reacrion_game')
def student_reacrion_game():
    return render_template('student/reacrion_game.html')

@app.route('/student/interview_wish')
def student_interview_wish():
    return render_template('student/interview_wish.html')

@app.route('/student/lucky_spin')
def student_lucky_spin():
    return render_template('student/lucky_spin.html')

@app.route('/student/choose_song')
def student_choose_song():
    return render_template('student/choose_song.html')

@app.route('/student/dinner_lucky_spin')
def student_dinner_lucky_spin():
    return render_template('student/dinner_lucky_spin.html')

@app.route('/student/snake_game')
def student_snake_game():
    return render_template('student/snake_game.html')

@app.route('/student/HSOS')
def student_HSOS():
    return render_template('student/HSOS.html')

@app.route('/student/HSOS_1')
def student_HSOS_1():
    return render_template('student/HSOS_1.html')

@app.route('/student/meeting')
def student_meeting():
    return render_template('student/meeting.html')


@app.route('/student/regulations.html')
def student_regulations():
    return render_template('student/regulations.html')



# ==========================================
# 第 5 區：學生服務中心(處理用路由)
# ==========================================

@app.route('/student/main_comment/submit', methods=['POST'])
def student_main_comment_submit():
    time_now = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")
    target_name = "海山高中" if request.form.get('q1') == '1' else "班聯會"
    try:
        db.session.add(MainComment(
            date=time_now,
            to=target_name,
            comment=request.form.get('feedback', ''),
            check=0
        ))
        db.session.commit()
        add_count('main_comment', increment=True)
    except Exception as e:
        db.session.rollback()
        db.session.add(MisTake(
            area = "匿名意見回饋",
            message = f"電腦給的錯誤為：{str(e)}"
        ))
        db.session.commit()
        return """<script>alert("系統發生一點小錯誤，儲存失敗，我們將盡速修改，請稍後再試！"); window.location.href = "/student/main_comment";</script>"""
    return """<script>alert("資料已送出！"); window.location.href = "/student/main";</script>"""

@app.route('/student/bandai/submit', methods=['POST'])
def student_bandai_submit():
    time_now = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")
    try:
        db.session.add(BandaiSignup(
            date = time_now,
            seat = request.form.get('student_id'),
            name = request.form.get('student_name'),
            check = 0
        ))
        db.session.commit()
        add_count('bandai_signup', increment=True)
    except Exception as e:
        db.session.rollback()
        db.session.add(MisTake(
            area = "報名列席班代大會",
            message = f"電腦給的錯誤為：{str(e)}"
        ))
        db.session.commit()
        return """<script>alert("系統發生一點小錯誤，儲存失敗，我們將盡速修改，請稍後再試！"); window.location.href = "/student/bandai";</script>"""
    return """<script>alert("資料已送出！"); window.location.href = "/student/main";</script>"""

@app.route('/student/interview_wish/submit', methods=['POST'])
def student_interview_wish_submit():
    time_now = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")
    user_choose = request.form.get('wish_type')
    user_interviewee = request.form.get('interviewee') # 抓到使用者打的受訪人姓名
    user_reason = request.form.get('reason')           # 抓到使用者打的原因
    user_topic_1 = request.form.get('topic_1')
    user_topic_2 = request.form.get('topic_2')
    user_topic = ""
    if user_choose == "許願受訪人":
        user_topic = user_topic_1
    else:
        user_topic = user_topic_2
    if user_interviewee == "":
        user_interviewee = None
    if user_reason == "":
        user_reason = None
    try:
        db.session.add(InterviewWish(
            date = time_now,
            choose=user_choose,
            interviewee=user_interviewee,
            reason=user_reason,
            topic=user_topic if user_topic else '無'
        ))
        db.session.commit()
        add_count('interview_wish', increment=True)
    except Exception as e:
        db.session.rollback()
        db.session.add(MisTake(
            area = "HSOS名人訪談許願區",
            message = f"電腦給的錯誤為：{str(e)}"
        ))
        db.session.commit()
        return """<script>alert("系統發生一點小錯誤，儲存失敗，我們將盡速修改，請再試一次！"); window.location.href = "/student/interview_wish";</script>"""
    return """<script>alert("資料已送出！"); window.location.href = "/student/main";</script>"""

@app.route('/student/choose_song/submit', methods=['POST'])
def student_choose_song_submit():
    time_now = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")
    user_song_name = request.form.get('song_name')
    user_artist = request.form.get('artist')
    user_link = request.form.get('yt_link')
    if user_artist == "":
        user_artist = None
    if user_link == "":
        user_link = None
    try:
        db.session.add(ChooseSong(
            date = time_now,
            song_name = user_song_name if user_song_name else '無'  ,
            artist = user_artist,
            yt_link = user_link
        ))
        db.session.commit()
        add_count('choose_song', increment=True)
    except Exception as e:
        db.session.rollback()
        db.session.add(MisTake(
            area = "廣播點歌許願區",
            message = f"電腦給的錯誤為：{str(e)}"
        ))
        db.session.commit()
        return """<script>alert("系統發生一點小錯誤，儲存失敗，我們將盡速修改，請再試一次！"); window.location.href = "/student/choose_song";</script>"""
    return """<script>alert("資料已送出！"); window.location.href = "/student/main";</script>"""

@app.route('/student/teacher_search/submit', methods=['POST'])
def student_teacher_search_submit():
    mode, keyword, subject = request.form.get('search_mode'), request.form.get('teacher_name', '').strip(), request.form.get('subject')
    teachers = []
    if os.path.exists(CSV_PATH_TEACHERS):
        with open(CSV_PATH_TEACHERS, mode='r', encoding='utf-8-sig') as f:
            teachers = list(csv.DictReader(f))
    else:
        db.session.add(MisTake(
            area = "教師辦公室查詢",
            message = f"這代表沒有路徑或 CSV 檔案遺失"
        ))
        db.session.commit()

    results = []
    if mode == 'by_name' and keyword:
        results = sorted([t for t in teachers if fuzz.partial_ratio(keyword, t['老師名稱']) >= 40], key=lambda x: fuzz.partial_ratio(keyword, x['老師名稱']), reverse=True)
    elif mode == 'by_filter' and subject:
        results = [t for t in teachers if t['任教科目'] == subject]

    return render_template('/student/teacher_result.html', results=results, count=len(results))

@app.route('/student/score_calculation/submit', methods=['POST'])
def student_score_calculation_submit():
    grade = request.form.get('grade')
    score1 = request.form.get('1')
    score2 = request.form.get('2')
    score3 = request.form.get('3')
    score4 = request.form.get('4')
    score5 = request.form.get('5')
    score6 = request.form.get('6')
    score7 = request.form.get('7')
    score8 = request.form.get('8')
    score9 = request.form.get('9')
    score10 = request.form.get('10')

    score = [score1,score2,score3,score4,score5,score6,score7,score8,score9,score10]
    subjects = ["國文","英文","數學","地理","歷史","公民","物理","化學","生物","地科"]
    weighted1 = [4,4,4,2,2,2,2,2,2,2]
    weighted3 = [0,4,4,0,0,0,3,3,2,1]
    use_weighted = []

    if grade == '1':
        use_weighted = weighted1
    else:
        use_weighted = weighted3

    use_subjects = []
    total = 0
    total_score = 0
    total_weighted = 0

    for i, s in enumerate(score):
        if not s.strip():
            continue
        else:
            try:
                subject_score = float(score[i])
            except ValueError:
                continue

            use_subjects.append(f"{subjects[i]}(加權{use_weighted[i]}) ： {subject_score}")
            total += use_weighted[i]
            total_score += subject_score * use_weighted[i]
            total_weighted += use_weighted[i]

    total *= 100
    if total_weighted == 0:
        avg = 0
    else:
        avg = total_score / total_weighted

    return render_template('student/score_result.html',subjects = use_subjects,total_score = total_score,total_weighted = total_weighted,total = total,avg = avg)

# ==========================================
# 第 6 區：班聯會管理頁面(登入和純路由)
# ==========================================

# 幾a幾b
@app.route('/abab')
def abab():
    return render_template('/abab.html')

@app.route('/abab_2')
def abab_2():
    return render_template('/abab_2.html')


DEPARTMENTS = {
    "chair": "主席",
    "academics": "學權",
    "pr": "公關",
    "sec": "秘書",
    "general": "總務",
    "it": "資訊",
    "manager":"管理者"
}

@app.route('/manager/student/login')
def manager_student_login():
    return render_template('manager/student/login.html')

@app.route('/manager/student/login/submit', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def manager_student_login_submit():
    if request.method == 'POST':
        pwd = request.form.get('password')
        password_map = {
            os.getenv('ADMIN_chair_PASSWORD', 'fallback_chair'): 'chair',
            os.getenv('ADMIN_academics_PASSWORD', 'fallback_academics'): 'academics',
            os.getenv('ADMIN_pr_PASSWORD', 'fallback_pr'): 'pr',
            os.getenv('ADMIN_sec_PASSWORD', 'fallback_sec'): 'sec',
            os.getenv('ADMIN_general_PASSWORD', 'fallback_general'): 'general',
            os.getenv('ADMIN_it_PASSWORD', 'fallback_it'): 'it',
            os.getenv('ADMIN_manager_PASSWORD', 'fallback_manager'): 'manager'
        }
        if pwd in password_map:
            matched_dept = password_map[pwd]
            session['manager'] = matched_dept
            return redirect(url_for('manager_student_student_interface'))
        else:
            return """<script>alert("密碼錯誤！"); window.location.href = "/manager/student/login";</script>"""
    return render_template('student/main.html')

# 純密碼
def manager_student_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('manager'):
            return """<script>alert("權限不足，請先登入！"); window.location.href = "/manager/student/login";</script>"""
        return f(*args, **kwargs)
    return decorated_function

# 分組
def manager_dept_required(allowed_depts):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            current_dept = session.get('manager')
            if current_dept not in allowed_depts:
                return """<script>alert("您所屬的組別無權限存取此功能！"); window.location.href = "/manager/student/student_interface";</script>"""
            return f(*args, **kwargs)
        return decorated_function
    return decorator

@app.route('/manager/student/student_interface')
@manager_student_required
def manager_student_student_interface():
    current_dept = session.get('manager')

    dept_ch = DEPARTMENTS.get(current_dept, "未知組別")

    # choose_song_count = ChooseSong.query.filter(ChooseSong.check != 1).count()
    choose_song_count = 0
    main_comment_count = MainComment.query.filter(MainComment.check != 1).count()
    line_message_count = LineMessage.query.filter(LineMessage.check != 1).count()
    error_count = MisTake.query.filter(MisTake.check != 1).count()

    return render_template(
        'manager/student/student_interface.html',
        dept=current_dept,
        dept_ch=dept_ch,
        choose_song_count=choose_song_count,
        main_comment_count=main_comment_count,
        line_message_count=line_message_count,
        error_count=error_count
    )

@app.route('/manager/student/line_message')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager'])
def manager_student_line_message():
    all_message = LineMessage.query.all()
    current_count = add_count('line_message')
    return render_template('manager/student/line_message.html', all_message=all_message, count=current_count)

@app.route('/manager/student/menu_comment')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager', 'academics', 'sec'])
def manager_student_menu_comment():
    return render_template('manager/student/menu_comment.html')

@app.route('/manager/student/interview_wish')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager', 'pr', 'sec'])
def manager_student_interview_wish():
    all_message = InterviewWish.query.all()
    current_count = add_count('interview_wish')
    return render_template('manager/student/interview_wish.html', all_message=all_message, count=current_count)

@app.route('/manager/student/choose_song')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager', 'sec'])
def manager_student_choose_song():
    all_message = ChooseSong.query.all()
    current_count = add_count('choose_song')
    return render_template('manager/student/choose_song.html', all_message=all_message, count=current_count)

@app.route('/manager/student/need_change')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager'])
def manager_student_need_change():
    return render_template('manager/student/need_change.html')

@app.route('/manager/student/error')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager'])
def manager_student_error():
    all_mistake = MisTake.query.all()
    return render_template('manager/student/error.html', all_mistake=all_mistake)

@app.route('/manager/student/main_comment')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager', 'academics', 'sec'])
def manager_student_main_comment():
    all_comment = MainComment.query.all()
    current_count = add_count('main_comment')
    return render_template('manager/student/main_comment.html', all_comment=all_comment, count=current_count)

@app.route('/manager/student/bandai_signup')
@manager_student_required
@manager_dept_required(['chair', 'it', 'manager', 'academics'])
def manager_student_bandai_signup():
    all_data = BandaiSignup.query.all()
    current_count = add_count('bandai_signup')
    return render_template('manager/student/bandai_signup.html', all_data=all_data, count=current_count)


# ==========================================
# 第 7 區：班聯會管理頁面(處理用路由)
# ==========================================

# 確認是否完成的[打勾]

@app.route('/update_check/<int:comment_id>', methods=['POST'])
@manager_student_required
def update_check(comment_id):
    comment = MainComment.query.get(comment_id)
    if comment:
        try:
            new_check = request.json.get('check')
            comment.check = new_check
            db.session.commit()
            return jsonify({"status": "success"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"status": "error", "message": "資料庫忙碌中，再試一次"}), 500
    return jsonify({"status": "error"}), 404

@app.route('/update_line_check/<int:msg_id>', methods=['POST'])
@manager_student_required
def update_line_check(msg_id):
    msg = LineMessage.query.get(msg_id)
    if msg:
        try:
            new_check = request.json.get('check')
            msg.check = new_check
            db.session.commit()
            return jsonify({"status": "success"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"status": "error", "message": "資料庫忙碌中，再試一次"}), 500
    return jsonify({"status": "error"}), 404

@app.route('/update_bandai_check/<int:signup_id>', methods=['POST'])
@manager_student_required
def update_bandai_check(signup_id):
    signup = BandaiSignup.query.get(signup_id)
    if signup:
        try:
            new_check = request.json.get('check')
            signup.check = new_check
            db.session.commit()
            return jsonify({"status": "success"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"status": "error", "message": "資料庫忙碌中，再試一次"}), 500
    return jsonify({"status": "error"}), 404

@app.route('/update_error_check/<int:signup_id>', methods=['POST'])
@manager_student_required
def update_error_check(signup_id):
    signup = MisTake.query.get(signup_id)
    if signup:
        try:
            new_check = request.json.get('check')
            signup.check = new_check
            db.session.commit()
            return jsonify({"status": "success"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"status": "error", "message": "資料庫忙碌中，再試一次"}), 500
    return jsonify({"status": "error"}), 404

@app.route('/update_choose_song/<int:signup_id>', methods=['POST'])
@manager_student_required
def update_choose_song(signup_id):
    signup = ChooseSong.query.get(signup_id)
    if signup:
        try:
            new_check = request.json.get('check')
            signup.check = new_check
            db.session.commit()
            return jsonify({"status": "success"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"status": "error", "message": "資料庫忙碌中，再試一次"}), 500
    return jsonify({"status": "error"}), 404

@app.route('/update_interview_wish/<int:signup_id>', methods=['POST'])
@manager_student_required
def update_interview_wish(signup_id):
    signup = InterviewWish.query.get(signup_id)
    if signup:
        try:
            new_check = request.json.get('check')
            signup.check = new_check
            db.session.commit()
            return jsonify({"status": "success"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"status": "error", "message": "資料庫忙碌中，再試一次"}), 500
    return jsonify({"status": "error"}), 404


# 刪除意見

@app.route('/manager/student/bandai_signup/delete', methods=['POST'])
@manager_student_required
@manager_dept_required(['it', 'manager'])
def manager_student_bandai_signup_delete():
    input_id = request.form.get('id')
    if not input_id or not input_id.isdigit():
        return """<script>alert("❌ 請輸入正確的數字編號！"); window.location.href = '/manager/student/bandai_signup';</script>"""
    item = BandaiSignup.query.get(int(input_id))
    if item:
        try:
            db.session.delete(item)
            db.session.commit()
            return """<script>alert("🗑️ 資料已成功刪除！"); window.location.href = '/manager/student/bandai_signup';</script>"""
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 資料庫忙碌中，請再試一次！"); window.location.href = document.referrer;</script>"""
    return """<script>alert("找不到這筆資料。"); window.location.href = '/manager/student/bandai_signup';</script>"""

@app.route('/manager/student/main_comment/delete', methods=['POST'])
@manager_student_required
@manager_dept_required(['it', 'manager'])
def manager_student_main_comment_delete():
    input_id = request.form.get('id')
    if not input_id or not input_id.isdigit():
        return """<script>alert("❌ 請輸入正確的數字編號！"); window.location.href = '/manager/student/main_comment';</script>"""
    item = MainComment.query.get(int(input_id))
    if item:
        try:
            db.session.delete(item)
            db.session.commit()
            return """<script>alert("🗑️ 資料已成功刪除！"); window.location.href = '/manager/student/main_comment';</script>"""
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 資料庫忙碌中，請再試一次！"); window.location.href = document.referrer;</script>"""
    return """<script>alert("找不到這筆資料。"); window.location.href = '/manager/student/main_comment';</script>"""

@app.route('/manager/student/line_message/delete', methods=['POST'])
@manager_student_required
@manager_dept_required(['it', 'manager'])
def manager_student_line_message_delete():
    input_id = request.form.get('id')
    if not input_id or not input_id.isdigit():
        return """<script>alert("❌ 請輸入正確的數字編號！"); window.location.href = '/manager/student/line_message';</script>"""
    item = LineMessage.query.get(int(input_id))
    if item:
        try:
            db.session.delete(item)
            db.session.commit()
            return """<script>alert("🗑️ 資料已成功刪除！"); window.location.href = '/manager/student/line_message';</script>"""
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 資料庫忙碌中，請再試一次！"); window.location.href = document.referrer;</script>"""
    return """<script>alert("找不到這筆資料。"); window.location.href = '/manager/student/line_message';</script>"""

@app.route('/manager/student/error/delete', methods=['POST'])
@manager_student_required
@manager_dept_required(['it', 'manager'])
def manager_student_error_delete():
    input_id = request.form.get('id')
    if not input_id or not input_id.isdigit():
        return """<script>alert("❌ 請輸入正確的數字編號！"); window.location.href = '/manager/student/error';</script>"""
    item = MisTake.query.get(int(input_id))
    if item:
        try:
            db.session.delete(item)
            db.session.commit()
            return """<script>alert("🗑️ 資料已成功刪除！"); window.location.href = '/manager/student/error';</script>"""
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 資料庫忙碌中，請再試一次！"); window.location.href = document.referrer;</script>"""
    return """<script>alert("找不到這筆資料。"); window.location.href = '/manager/student/error';</script>"""

@app.route('/manager/student/choose_song/delete', methods=['POST'])
@manager_student_required
@manager_dept_required(['it', 'manager'])
def manager_student_choose_song_delete():
    input_id = request.form.get('id')
    if not input_id or not input_id.isdigit():
        return """<script>alert("❌ 請輸入正確的數字編號！"); window.location.href = '/manager/student/choose_song';</script>"""
    item = ChooseSong.query.get(int(input_id))
    if item:
        try:
            db.session.delete(item)
            db.session.commit()
            return """<script>alert("🗑️ 資料已成功刪除！"); window.location.href = '/manager/student/choose_song';</script>"""
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 資料庫忙碌中，請再試一次！"); window.location.href = document.referrer;</script>"""
    return """<script>alert("找不到這筆資料。"); window.location.href = '/manager/student/choose_song';</script>"""

@app.route('/manager/student/interview_wish/delete', methods=['POST'])
@manager_student_required
@manager_dept_required(['it', 'manager'])
def manager_student_interview_wish_delete():
    input_id = request.form.get('id')
    if not input_id or not input_id.isdigit():
        return """<script>alert("❌ 請輸入正確的數字編號！"); window.location.href = '/manager/student/interview_wish';</script>"""
    item = InterviewWish.query.get(int(input_id))
    if item:
        try:
            db.session.delete(item)
            db.session.commit()
            return """<script>alert("🗑️ 資料已成功刪除！"); window.location.href = '/manager/student/interview_wish';</script>"""
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 資料庫忙碌中，請再試一次！"); window.location.href = document.referrer;</script>"""
    return """<script>alert("找不到這筆資料。"); window.location.href = '/manager/student/interview_wish';</script>"""

# ==========================================
# 第 8 區：營養午餐意見回饋
# ==========================================

@app.route('/menu_comment/create')
def menu_comment_create():
    return render_template('menu_comment/create.html')

@app.route('/menu_comment/create/submit', methods=['POST'])
def menu_comment_create_submit():
    shop = request.form.get('shop')
    if shop not in ["上將", "宏遠", "統鮮"]:
        return """<script>alert("❌ 無效的廠商參數！"); window.location.href = '/menu_comment/create';</script>"""

    now = datetime.now(timezone(timedelta(hours=8)))
    dates = [f"{(now - timedelta(days=i)).month}/{(now - timedelta(days=i)).day}" for i in range(3)]


    menu_data = {}
    csv_path = os.path.join(BASE_DIR, 'menus', f"{now.year}年{now.month}月海山高中菜單_{shop}.csv")
    if os.path.exists(csv_path):
        with open(csv_path, mode='r', encoding='utf-8-sig') as f:
            for row in csv.DictReader(f):
                items = row.get('品項', '')
                menu_data[row.get('日期', '').strip()] = items.split('、') if items else []
    return render_template('menu_comment/form.html', shop=shop, menu_data=menu_data, today=dates[0], yesterday=dates[1], day_before=dates[2], is_officer=request.form.get('ys'), grade=request.form.get('q1'), class_num=request.form.get('class_num'))

@app.route('/menu_comment/form/submit', methods=['POST'])
def menu_comment_form_submit():
    data = request.form
    student_class = f"{data.get('grade')}{data.get('class_num')}" if data.get('is_officer') == 'yes' else "無"

    date = data.get('date')
    shop = data.get('shop')
    too_much_food = data.get('too_much_food', '未填寫')

    dish_names = request.form.getlist('dish_name')
    dish_comments = request.form.getlist('dish_comment')

    has_db_change = False

    db.session.add(DishComment(
        date=date,
        shop=shop,
        dish_name='【整日份量調查】',
        student_class=student_class,
        comment=f'午餐是否過多: {too_much_food}'
    ))
    add_count('dish_comment', increment=True)
    has_db_change = True

    for d_name, d_comment in zip(dish_names, dish_comments):
        d_comment_clean = d_comment.strip()

        if d_comment_clean:
            db.session.add(DishComment(
                date=date,
                shop=shop,
                dish_name=d_name,
                student_class=student_class,
                comment=d_comment_clean
            ))
            add_count('dish_comment', increment=True)

    if has_db_change:
        try:
            db.session.commit()
        except Exception as e:
            db.session.rollback()
            return """<script>alert("❌ 系統忙碌中，儲存失敗，請再試一次！"); window.location.href = '/menu_comment/create';</script>"""

    return """<script>alert("表單已成功送出！感謝你的建議。"); window.location.href='/student/main';</script>"""

# ==========================================
# 第 9 區：營養午餐管理後台
# ==========================================

@app.route('/manager/menu_comment/all_login')
def manager_menu_comment_all_login():
    return render_template('manager/menu_comment/all_login.html')

@app.route('/manager/menu_comment/all_login/submit', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def manager_menu_comment_all_login_submit():
    if request.method == 'POST':
        pwd = request.form.get('admin_password')

        manufacturer_password = os.getenv('Menu_comment_Manufacturer_PASSWORD', 'fallback_random_string')
        teacher_password = os.getenv('Menu_comment_Teacher_PASSWORD', 'fallback_random_string')

        password_map = {
            os.getenv('ADMIN_chair_PASSWORD', 'fallback_chair'): 'chair',
            os.getenv('ADMIN_academics_PASSWORD', 'fallback_academics'): 'academics',
            os.getenv('ADMIN_sec_PASSWORD', 'fallback_sec'): 'sec',
            os.getenv('ADMIN_it_PASSWORD', 'fallback_it'): 'it',
            os.getenv('ADMIN_manager_PASSWORD', 'fallback_manager'): 'manager'
        }

        if pwd == manufacturer_password:
            session.permanent = False
            session['all_food'] = True
            return redirect(url_for('manager_menu_comment_manager'))

        elif pwd == teacher_password:
            session.permanent = False
            session['all_food'] = True
            session['banlian_food'] = True
            return redirect(url_for('manager_menu_comment_manager'))

        elif pwd in password_map:
            matched_dept = password_map[pwd]
            session.permanent = False
            session['manager'] = matched_dept
            session['all_food'] = True
            session['banlian_food'] = True
            return redirect(url_for('manager_menu_comment_manager'))

        else:
            return """<script>alert("❌ 密碼錯誤！"); window.location.href = "/manager/menu_comment/all_login";</script>"""
    return render_template('student/main.html')

@app.route('/manager/menu_comment/banlian_login')
def manager_menu_comment_banlian_login():
    return render_template('manager/menu_comment/banlian_login.html')

@app.route('/manager/menu_comment/banlian_login/submit', methods=['GET', 'POST'])
@limiter.limit("5 per minute")
def manager_menu_comment_banlian_login_submit():
    if request.method == 'POST':
        pwd = request.form.get('admin_password')

        teacher_password = os.getenv('Menu_comment_Teacher_PASSWORD', 'fallback_random_string')

        password_map = {
            os.getenv('ADMIN_chair_PASSWORD', 'fallback_chair'): 'chair',
            os.getenv('ADMIN_academics_PASSWORD', 'fallback_academics'): 'academics',
            os.getenv('ADMIN_sec_PASSWORD', 'fallback_sec'): 'sec',
            os.getenv('ADMIN_it_PASSWORD', 'fallback_it'): 'it',
            os.getenv('ADMIN_manager_PASSWORD', 'fallback_manager'): 'manager'
        }

        if pwd == teacher_password:
            session.permanent = False
            session['banlian_food'] = True
            session['all_food'] = True
            return redirect(url_for('manager_menu_comment_all_comment'))

        elif pwd in password_map:
            matched_dept = password_map[pwd]
            session.permanent = False
            session['manager'] = matched_dept
            session['banlian_food'] = True
            session['all_food'] = True
            return redirect(url_for('manager_menu_comment_all_comment'))

        else:
            return """<script>alert("❌ 密碼錯誤！"); window.location.href = "/manager/menu_comment/banlian_login";</script>"""
    return render_template('student/main.html')

def manager_menu_comment_all_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        has_all_food = session.get('all_food')
        is_advanced_admin = session.get('manager') in ['chair','academics','sec','it', 'manager']

        if not (has_all_food or is_advanced_admin):
            return """<script>alert("權限不足，請先登入！"); window.location.href = "/manager/menu_comment/all_login";</script>"""
        return f(*args, **kwargs)
    return decorated_function

def manager_menu_comment_banlian_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        has_banlian_food = session.get('banlian_food')
        is_advanced_admin = session.get('manager') in ['chair','academics','sec','it', 'manager']

        if not (has_banlian_food or is_advanced_admin):
            return """<script>alert("權限不足，請先登入！"); window.location.href = "/manager/menu_comment/banlian_login";</script>"""
        return f(*args, **kwargs)
    return decorated_function

@app.route('/manager/menu_comment/manager', methods=['GET'])
@manager_menu_comment_all_required
def manager_menu_comment_manager():
    mode = request.args.get('mode', 'week')
    week = request.args.get('week')
    month = request.args.get('month')
    shop = request.args.get('shop')

    query = DishComment.query
    query = query.filter(DishComment.dish_name != '【整日份量調查】')

    if shop:
        query = query.filter(DishComment.shop == shop)

    if mode == 'week' and week:
        try:
            monday = datetime.strptime(f'{week}-1', "%G-W%V-%u")
            db_dates = [f"{(monday + timedelta(days=i)).month}/{(monday + timedelta(days=i)).day}" for i in range(5)]
            query = query.filter(DishComment.date.in_(db_dates))
        except Exception as e:
            # 移除 pass，改為印出錯誤，避免 Silent Fail
            print(f"日期解析錯誤: {e}")

    elif mode == 'month' and month:
        try:
            m_target = f"{int(month.split('-')[1])}/"
            query = query.filter(DishComment.date.like(f"{m_target}%"))
        except Exception as e:
            # 移除 pass，改為印出錯誤
            print(f"月份解析錯誤: {e}")

    all_comments = query.all() if shop and ((mode == 'week' and week) or (mode == 'month' and month)) else []
    all_comments.sort(key=lambda x: int(x.date.split('/')[0])*100 + int(x.date.split('/')[1]) if '/' in x.date else 0)

    structured_data = {}
    for c in all_comments:
        structured_data.setdefault(c.date, {}).setdefault(c.dish_name, []).append(c)

    res = make_response(render_template('manager/menu_comment/manager.html',
                                        structured_data=structured_data,
                                        selected_mode=mode,
                                        selected_week=week,
                                        selected_month=month,
                                        selected_shop=shop))
    res.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    return res

@app.route('/manager/menu_comment/class_comment')
@manager_menu_comment_banlian_required
def manager_menu_comment_class_comment():
    mode, week, month = request.args.get('mode', 'week'), request.args.get('week'), request.args.get('month')
    db_dates, expected_days = [], 0

    if mode == 'week' and week:
        try:
            monday = datetime.strptime(f'{week}-1', "%G-W%V-%u")
            db_dates = [f"{(monday + timedelta(days=i)).month}/{(monday + timedelta(days=i)).day}" for i in range(5)]
            expected_days = 5
        except Exception as e:
            print(f"日期解析錯誤: {e}")

    elif mode == 'month' and month:
        try:
            y, m = map(int, month.split('-'))
            db_dates = [f"{m}/{day}" for day in range(1, calendar.monthrange(y, m)[1] + 1) if datetime(y, m, day).weekday() < 5]
            expected_days = len(db_dates)
        except Exception as e:
            print(f"月份解析錯誤: {e}")

    all_classes = [f"{g}{c:02d}" for g in (1, 2, 3) for c in range(1, 11)]
    class_counts = {c: 0 for c in all_classes}

    if db_dates:
        class_daily = {c: set() for c in all_classes}
        for c in DishComment.query.filter(DishComment.date.in_(db_dates)).all():
            if c.student_class in class_daily: class_daily[c.student_class].add(c.date)
        for c in all_classes: class_counts[c] = len(class_daily[c])

    lagging_classes = [c for c, count in class_counts.items() if count < int(expected_days / 2)] if expected_days > 0 else []
    return render_template('manager/menu_comment/class_comment.html', selected_mode=mode, selected_week=week, selected_month=month, expected_days=expected_days, class_counts=class_counts, lagging_classes=lagging_classes, all_classes=all_classes)

@app.route('/manager/menu_comment/all_comment')
@manager_menu_comment_banlian_required
def manager_menu_comment_all_comment():
    all_comments = DishComment.query.filter(DishComment.dish_name != '【整日份量調查】').all()
    current_count = add_count('dish_comment')
    all_comments.sort(key=lambda x: int(x.date.split('/')[0])*100 + int(x.date.split('/')[1]) if '/' in x.date else 0, reverse=True)
    res = make_response(render_template('manager/menu_comment/all_comment.html', all_comments=all_comments,count = current_count))
    res.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    return res

@app.route('/manager/menu_comment/delete_comment/<int:comment_id>', methods=['POST'])
@manager_menu_comment_banlian_required
def manager_menu_comment_delete_comment(comment_id):
    comment = DishComment.query.get(comment_id)
    if comment:
        db.session.delete(comment)
        db.session.commit()
        return """<script>alert("🗑️ 留言已成功刪除！"); window.location.href = document.referrer || '/manager/menu_comment/manager';</script>"""
    return """<script>alert("找不到這筆留言，可能已經被刪除了。"); window.location.href = document.referrer || '/manager/menu_comment/manager';</script>"""

# ==========================================
# 第 10 區：帳號登入與權限驗證專區
# ==========================================

oauth = OAuth(app)
google = oauth.register(
    name='google',
    client_id=os.getenv('GOOGLE_CLIENT_ID'),
    client_secret=os.getenv('GOOGLE_CLIENT_SECRET'),
    server_metadata_url='https://accounts.google.com/.well-known/openid-configuration',
    client_kwargs={'scope': 'openid email profile'}
)

# 💡 這是我們自訂的「時間門禁警衛」裝飾器
def check_preorder_time(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        setting = OrderSetting.query.first()
        tz = timezone(timedelta(hours=8))
        now = datetime.now(tz)

        is_locked = True
        if setting and setting.start_time and setting.end_time:
            start = datetime.strptime(setting.start_time, '%Y-%m-%dT%H:%M').replace(tzinfo=tz)
            end = datetime.strptime(setting.end_time, '%Y-%m-%dT%H:%M').replace(tzinfo=tz)
            if start <= now <= end:
                is_locked = False

        # 如果時間不對，直接阻擋並彈回首頁
        if is_locked:
            return """<script>alert("目前非預購開放時間！系統已鎖定。"); window.location.href = '/student/main';</script>"""

        # 時間正確，才讓他繼續執行原本的網頁功能
        return f(*args, **kwargs)
    return decorated_function


@app.route('/order/google_login')
@check_preorder_time # 加上警衛
def order_google_login():
    return render_template('order/google_login.html')

@app.route('/order/login_student')
@check_preorder_time # 加上警衛
def order_login_student():
    return google.authorize_redirect(url_for('order_student_google_callback', _external=True), prompt='select_account')

@app.route('/order/student_google_callback')
@check_preorder_time # 加上警衛
def order_student_google_callback():
    try:
        token = google.authorize_access_token()
        user_info = google.get('https://www.googleapis.com/oauth2/v3/userinfo').json()
        if not user_info.get('email', '').endswith('@hshs.ntpc.edu.tw'):
            return """<script>alert("請使用學校信箱登入！"); window.location.href = "/order/google_login";</script>"""
        session['user-student'] = {'email': user_info['email'], 'name': user_info.get('name', ''), 'picture': user_info.get('picture', '')}

        user_login_id = user_info['email'].split('@')[0]
        existing_user = OrderUser.query.filter_by(login_id=user_login_id).first()
        if existing_user:
            session['order'] = session['user-student']
            return redirect(url_for('order_index_view'))

        return redirect(url_for('order_student_profile_setup'))
    except Exception as e:
        return f"登入出錯了：{str(e)}"

@app.route('/order/student_profile_setup')
@check_preorder_time # 加上警衛
def order_student_profile_setup():
    session.permanent = False
    session['order'] = session.get('user-student')
    if not session.get('user-student'):
        return """<script>alert("請登入學生身分");window.location.href = "/order/google_login";</script>"""
    return render_template('order/student_profile_setup.html', student=session.get('user-student'))

@app.route('/order/index', methods=['POST'])
@check_preorder_time # 加上警衛
def order_index():
    user = session.get('order')
    if not user:
        return """<script>alert("權限不足，請先登入！"); window.location.href = "/order/google_login";</script>"""

    choose = request.form.get('search_mode')
    grade = request.form.get('grade')
    class_num = request.form.get('class_num')
    seat_num = request.form.get('seat_num')

    if not grade or not class_num or not seat_num:
        return """<script>alert("請確認年級、班級、座號都已完整填寫！"); window.history.back();</script>"""

    is_student = True if choose == 'by_student' else False
    class_name = str(grade) + str(class_num)

    user_login_id = user['email'].split('@')[0]
    exist_user = OrderUser.query.filter_by(login_id=user_login_id).first()
    if not exist_user:
        db.session.add(OrderUser(
            is_student=is_student,
            login_id=user_login_id,
            class_name=class_name,
            seat_number=seat_num,
            student_name=user['name']
        ))
        db.session.commit()

    # 註：這裡改為導向 index_view 以保持頁面邏輯統一
    return redirect(url_for('order_index_view'))

# 購物車與商品主頁面
@app.route('/order/index_view')
@check_preorder_time # 加上警衛
def order_index_view():
    user = session.get('order')
    cart_items = []
    class_name = seat_number = name = None
    order_status = '購物車'

    if user:
        user_db = OrderUser.query.filter_by(login_id=user['email'].split('@')[0]).first()
        if user_db:
            class_name = user_db.class_name
            seat_number = user_db.seat_number
            name = user_db.student_name
            order = OrderMain.query.filter(
                OrderMain.user_id == user_db.id,
                OrderMain.status.in_(['購物車', '已送出'])
            ).order_by(OrderMain.order_id.desc()).first()
            if order:
                cart_items = order.items
                order_status = order.status

    products = OrderProduct.query.all()

    # 取出系統設定，傳給前端以顯示跑馬燈跟倒數計時器
    setting = OrderSetting.query.first()

    return render_template('order/index.html', products=products, cart_items=cart_items,
                            class_name=class_name, seat_number=seat_number, name=name,
                            order_status=order_status, setting=setting)

# 加入購物車 (防範學生駭客送假資料硬買)
@app.route('/order/add_to_cart', methods=['POST'])
@check_preorder_time # 加上警衛
def order_add_to_cart():
    user_data = session.get('order')
    if not user_data:
        return """<script>alert("請先登入！"); window.location.href = "/order/google_login";</script>"""

    user_login_id = user_data['email'].split('@')[0]
    user = OrderUser.query.filter_by(login_id=user_login_id).first()

    if not user:
        return """<script>alert("找不到您的班級座號資料，請重新登入並填寫一次！"); window.location.href = "/order/google_login";</script>"""

    submitted_order = OrderMain.query.filter_by(user_id=user.id, status='已送出').first()
    if submitted_order:
        return """<script>alert("您已送出預購，如需修改請先取消預購！"); window.location.href = '/order/index_view';</script>"""

    product_id = request.form.get('product_id')
    size = request.form.get('size')

    try:
        quantity = int(request.form.get('quantity', 1))
    except ValueError:
        quantity = 1

    if quantity < 1:
        quantity = 1
    elif quantity > 10:
        quantity = 10

    product = OrderProduct.query.get(product_id)
    if not product:
        return """<script>alert("找不到該商品！"); window.location.href = '/order/index_view';</script>"""

    valid_sizes = product.sizes.split(',')
    if size not in valid_sizes:
        return """<script>alert("選擇的尺寸無效，請勿竄改資料！"); window.location.href = '/order/index_view';</script>"""

    try:
        order = OrderMain.query.filter_by(user_id=user.id, status='購物車').first()
        if not order:
            order = OrderMain(user_id=user.id, status='購物車')
            db.session.add(order)
            db.session.commit()

        existing_item = OrderItem.query.filter_by(
            order_id=order.order_id,
            product_id=product_id,
            item_size=size
        ).first()

        if existing_item:
            existing_item.quantity += quantity
        else:
            new_item = OrderItem(order_id=order.order_id, product_id=product_id, item_size=size, quantity=quantity)
            db.session.add(new_item)

        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return f"""<script>alert("加入購物車失敗，請聯絡管理員。錯誤：{str(e)}\\n"); window.location.href = '/order/index_view';</script>"""
    return """<script>alert("✅ 已加入購物車！"); window.location.href = '/order/index_view';</script>"""

# 修改數量
@app.route('/order/cart/update_quantity', methods=['POST'])
@check_preorder_time # 加上警衛
def order_cart_update_quantity():
    user_data = session.get('order')
    if not user_data:
        return """<script>alert("請先登入！"); window.location.href = "/order/google_login";</script>"""

    user = OrderUser.query.filter_by(login_id=user_data['email'].split('@')[0]).first()
    if not user:
        return """<script>alert("找不到您的班級座號資料，請重新登入並填寫一次！"); window.location.href = "/order/google_login";</script>"""
    item_id = request.form.get('item_id')
    try:
        new_quantity = int(request.form.get('quantity', 1))
    except (TypeError, ValueError):
        new_quantity = 1

    # ===== 新增防護：限制數量上限 =====
    if new_quantity < 1:
        new_quantity = 1
    elif new_quantity > 10:
        new_quantity = 10  # 強制將惡意修改的大數字壓回 10
    # ===============================

    item = OrderItem.query.get(item_id)
    if not item or item.order.user_id != user.id or item.order.status != '購物車':
        return """<script>alert("找不到這筆項目，或目前無法修改。"); window.location.href = '/order/index_view';</script>"""

    item.quantity = new_quantity
    db.session.commit()

    return """<script>window.location.href = '/order/index_view';</script>"""


# 刪除購物車項目
@app.route('/order/cart/delete_item', methods=['POST'])
@check_preorder_time # 加上警衛
def order_cart_delete_item():
    user_data = session.get('order')
    if not user_data:
        return """<script>alert("請先登入！"); window.location.href = "/order/google_login";</script>"""

    user = OrderUser.query.filter_by(login_id=user_data['email'].split('@')[0]).first()
    if not user:
        return """<script>alert("找不到您的班級座號資料，請重新登入並填寫一次！"); window.location.href = "/order/google_login";</script>"""
    item_id = request.form.get('item_id')

    item = OrderItem.query.get(item_id)
    if not item or item.order.user_id != user.id or item.order.status != '購物車':
        return """<script>alert("找不到這筆項目，或目前無法修改。"); window.location.href = '/order/index_view';</script>"""

    db.session.delete(item)
    db.session.commit()

    return """<script>window.location.href = '/order/index_view';</script>"""

# 送出預購
@app.route('/order/submit_order', methods=['POST'])
@check_preorder_time # 加上警衛
def order_submit_order():
    user_data = session.get('order')
    if not user_data:
        return """<script>alert("請先登入！"); window.location.href = "/order/google_login";</script>"""

    user = OrderUser.query.filter_by(login_id=user_data['email'].split('@')[0]).first()
    if not user:
        return """<script>alert("找不到您的班級座號資料，請重新登入並填寫一次！"); window.location.href = "/order/google_login";</script>"""

    if request.form.get('agree') != 'yes':
        return """<script>alert("請先勾選同意條款再送出！"); window.location.href = '/order/index_view';</script>"""

    order = OrderMain.query.filter_by(user_id=user.id, status='購物車').first()
    if not order or not order.items:
        return """<script>alert("購物車是空的，無法送出預購！"); window.location.href = '/order/index_view';</script>"""

    for item in order.items:
        if item.quantity > 10:
            return f"""<script>alert("【{item.product.name} ({item.item_size})】數量超過 10 件上限！請先修改購物車數量再送出。"); window.location.href = '/order/index_view';</script>"""

    order.status = '已送出'
    db.session.commit()

    return """<script>alert("🎉 預購已送出！"); window.location.href = '/order/index_view';</script>"""

# 取消送出預購
@app.route('/order/cancel_order', methods=['POST'])
@check_preorder_time # 加上警衛
def order_cancel_order():
    user_data = session.get('order')
    if not user_data:
        return """<script>alert("請先登入！"); window.location.href = "/order/google_login";</script>"""

    user = OrderUser.query.filter_by(login_id=user_data['email'].split('@')[0]).first()
    if not user:
        return """<script>alert("找不到您的班級座號資料，請重新登入並填寫一次！"); window.location.href = "/order/google_login";</script>"""

    order = OrderMain.query.filter_by(user_id=user.id, status='已送出').first()
    if order:
        order.status = '購物車'
        db.session.commit()

    return """<script>alert("已取消預購，購物車已可重新編輯。"); window.location.href = '/order/index_view';</script>"""

# ==========================================
# 第 11 區：校慶紀念服後台
# ==========================================

def get_excel_suffix():
    """
    例句：這是用來判斷現在時間的函式。
    如果現在還沒超過截止時間，或是根本還沒設定時間，就會回傳『非最終版』的字樣，加在 Excel 檔名後面。
    """
    setting = OrderSetting.query.first()
    if not setting or not setting.end_time:
        return "_非最終版"

    # 將字串轉換為時間物件並進行比對
    tz = timezone(timedelta(hours=8))
    now = datetime.now(tz)
    end_time = datetime.strptime(setting.end_time, '%Y-%m-%dT%H:%M').replace(tzinfo=tz)

    if now < end_time:
        return "_非最終版"
    return ""

@app.route('/manager/order/update_settings', methods=['POST'])
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_update_settings():
    setting = OrderSetting.query.first()
    setting.start_time = request.form.get('start_time')
    setting.end_time = request.form.get('end_time')
    # 這裡不再更新廣播了，交給另一個專屬路由處理
    db.session.commit()
    return """<script>alert("時間設定已啟動！"); window.location.href = '/manager/order/admin_dashboard';</script>"""

@app.route('/manager/order/update_marquee', methods=['POST'])
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_update_marquee():
    setting = OrderSetting.query.first()
    if not setting:
        setting = OrderSetting()
        db.session.add(setting)
    setting.marquee = request.form.get('marquee')
    db.session.commit()
    return """<script>alert("廣播內容已更新！"); window.location.href = '/manager/order/admin_dashboard';</script>"""

@app.route('/manager/order/hard_reset', methods=['POST'])
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_hard_reset():
    pwd = request.form.get('super_password')
    # 從環境變數讀取最高密碼，如果沒有設定，預設為 'HSHS_RESET_999'
    if pwd != os.getenv('SUPER_ADMIN_PASSWORD', 'HSHS_RESET_999'):
        return """<script>alert("❌ 最高密碼錯誤！"); window.location.href = '/manager/order/admin_dashboard';</script>"""

    try:
        # 1. 自動備份：建立專屬備份資料夾
        backup_dir = os.path.join(BASE_DIR, 'backups', datetime.now().strftime('重置備份_%Y%m%d_%H%M%S'))
        os.makedirs(backup_dir, exist_ok=True)

        # 將所有 Excel 產生的結果直接寫入資料夾
        files_to_backup = {
            'part1_各班分頁報表(班代用).xlsx': manager_order_export_excel,
            'part1_各班收費報表(班聯用).xlsx': manager_order_export_excel_class_collection,
            'part2_廠商訂購總表.xlsx': manager_order_export_excel_manufacturer,
            'part3_個人發放衣服對應表(班代用).xlsx': manager_order_export_excel_student_delivery,
            'part3_各班size總額(班聯用).xlsx': manager_order_export_excel_class_size,
            'other_全校總表.xlsx': manager_order_export_excel_all
        }

        for filename, func in files_to_backup.items():
            resp = func()

            # 👇 這裡就是修正的地方：解除穿透模式，安全讀取二進位資料 👇
            resp.direct_passthrough = False

            with open(os.path.join(backup_dir, filename), 'wb') as f:
                f.write(resp.get_data())

        # 2. 清除所有資料庫紀錄
        db.session.query(OrderItem).delete()
        db.session.query(OrderMain).delete()
        db.session.query(OrderUser).delete()
        db.session.query(OrderProduct).delete()

        # 解除時間與廣播鎖定
        setting = OrderSetting.query.first()
        if setting:
            setting.start_time = None
            setting.end_time = None
            setting.marquee = None
        db.session.commit()

        return """<script>alert("✅ 系統已重置完成！備份檔案已存入主機。"); window.location.href = '/manager/order/admin_dashboard';</script>"""
    except Exception as e:
        db.session.rollback()
        return f"""<script>alert("發生錯誤：{str(e)}"); window.history.back();</script>"""


@app.route('/manager/order/delete_product/<int:product_id>', methods=['POST'])
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_delete_product(product_id):

    setting = OrderSetting.query.first()
    tz = timezone(timedelta(hours=8))
    now = datetime.now(tz)
    if setting and setting.start_time:
        start = datetime.strptime(setting.start_time, '%Y-%m-%dT%H:%M').replace(tzinfo=tz)
        if now >= start:
            return """<script>alert("預購已啟動，嚴禁刪除商品！"); window.location.href = '/manager/order/product_list';</script>"""


    product = OrderProduct.query.get(product_id)
    if product:
        existing_items = OrderItem.query.filter_by(product_id=product.id).first()
        if existing_items:
            return """<script>alert("❌ 無法刪除！..."); window.location.href = '/manager/order/product_list';</script>"""

        # ▼ 新增以下這段：清除主機上的實際圖片檔案 ▼
        if product.images:
            # 1. 把資料庫裡以逗號分隔的網址字串，切成清單 (List)
            image_list = product.images.split(',')

            # 2. 針對清單裡的每一張圖片網址，逐一處理
            for img_url in image_list:
                if img_url:  # 確保網址不是空的
                    # 3. img_url 目前長這樣： "/static/order/xxx.png"
                    # 將其與 BASE_DIR ('./mysite') 組合，轉為作業系統看得懂的真實路徑
                    # 例如變成： "./mysite/static/order/xxx.png"
                    file_path = BASE_DIR + img_url

                    # 4. 透過 os.path 檢查這個檔案是否真的存在於硬碟裡
                    if os.path.exists(file_path):
                        # 5. 確認存在後，呼叫 os.remove 將其實體刪除
                        os.remove(file_path)
        # ▲ 新增結束 ▲

        db.session.delete(product)
        db.session.commit()
    return """<script>alert("商品已刪除！"); window.location.href = '/manager/order/product_list';</script>"""

# ------------------------------------------------------------------------------------------------

@app.route('/manager/order/admin_dashboard')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_admin_dashboard():
    orders = OrderMain.query.filter_by(status='已送出').all()
    unique_users = set(order.user_id for order in orders)
    total_users = len(unique_users)
    total_items = sum(item.quantity for order in orders for item in order.items)
    total_money = sum(item.product.price * item.quantity for order in orders for item in order.items)

    setting = OrderSetting.query.first()
    tz = timezone(timedelta(hours=8))
    now = datetime.now(tz)

    is_locked_excel = True
    if setting and setting.end_time:
        try:
            end = datetime.strptime(setting.end_time, '%Y-%m-%dT%H:%M').replace(tzinfo=tz)
            if now >= end:
                is_locked_excel = False
        except ValueError:
            pass

    return render_template('manager/order/admin_dashboard.html',
                           total_users=total_users,
                           total_items=total_items,
                           total_money=total_money,
                           setting=setting,
                           is_locked_excel=is_locked_excel)

@app.route('/manager/order/product_add')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_product_add():
    return render_template('manager/order/product_add.html')

@app.route('/manager/order/product_add/submit', methods=['POST'])
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_product_add_submit():
    # 🚨 新增防護：鎖定期間後端拒絕新增
    setting = OrderSetting.query.first()
    tz = timezone(timedelta(hours=8))
    now = datetime.now(tz)
    if setting and setting.start_time:
        start = datetime.strptime(setting.start_time, '%Y-%m-%dT%H:%M').replace(tzinfo=tz)
        if now >= start:
            return """<script>alert("預購已啟動，嚴禁新增商品！"); window.location.href = '/manager/order/admin_dashboard';</script>"""

    # ... (下方保留你原本的 product_name 等接收與存檔邏輯) ...
    product_name = request.form.get('product_name')
    product_price = request.form.get('product_price')
    sizes = request.form.getlist('sizes')
    sizes_str = ",".join(sizes)
    uploaded_files = request.files.getlist('product_images')
    image_paths = []
    upload_folder = os.path.join(BASE_DIR, 'static', 'order')
    if not os.path.exists(upload_folder):
        os.makedirs(upload_folder)
    for file in uploaded_files:
        if file.filename != '':
            filename = secure_filename(file.filename)
            file_path = os.path.join(upload_folder, filename)
            file.save(file_path)
            image_paths.append(f"/static/order/{filename}")
    images_str = ",".join(image_paths)
    new_product = OrderProduct(name=product_name, price=product_price, images=images_str, sizes=sizes_str)
    db.session.add(new_product)
    db.session.commit()
    return """<script>alert("上架成功！"); window.location.href = '/manager/order/product_list';</script>"""





@app.route('/manager/order/product_list')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_product_list():
    products = OrderProduct.query.all()
    return render_template('manager/order/product_list.html', products=products)


# 匯出 excel-1 班級座號姓名全校版
@app.route('/manager/order/export_excel')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_export_excel():
    products = OrderProduct.query.all()

    orders = db.session.query(OrderMain).join(OrderUser)\
             .filter(OrderMain.status == '已送出')\
             .order_by(OrderUser.class_name, OrderUser.seat_number).all()

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    fill_grey = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    orders_by_class = groupby(orders, key=lambda o: o.user.class_name)

    for class_name, class_orders in orders_by_class:
        class_orders_list = list(class_orders)
        class_total_money = 0
        current_row = 1

        ws = wb.create_sheet(title=f"{class_name}班")

        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 85
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.print_title_cols = 'A:D'
        ws.print_title_rows = '1:2'
        ws.freeze_panes = 'E3'

        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 6
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 12

        ws.cell(row=current_row, column=1, value=f"--- {class_name} 班 訂購表 ---").font = Font(bold=True, size=14)

        col_idx = 5
        size_accum = 0
        product_col_map = {}

        for idx, p in enumerate(products):
            sizes = p.sizes.split(',')
            ws.merge_cells(start_row=current_row, start_column=col_idx, end_row=current_row, end_column=col_idx + len(sizes) - 1)

            cell = ws.cell(row=current_row, column=col_idx, value=f"{p.name} ({p.price}元)")
            cell.alignment = align_center
            cell.border = border

            product_col_map[p.id] = {'start': col_idx, 'sizes': sizes, 'color_idx': idx}

            for s_col in range(col_idx, col_idx + len(sizes)):
                ws.column_dimensions[openpyxl.utils.get_column_letter(s_col)].width = 6

            col_idx += len(sizes)
            size_accum += len(sizes)

            if size_accum > 18 and idx != len(products) - 1:
                ws.col_breaks.append(Break(id=col_idx - 1))
                size_accum = 0

        money_col, sign_col, check_col = col_idx, col_idx + 1, col_idx + 2

        ws.column_dimensions[openpyxl.utils.get_column_letter(money_col)].width = 10
        ws.column_dimensions[openpyxl.utils.get_column_letter(sign_col)].width = 16
        ws.column_dimensions[openpyxl.utils.get_column_letter(check_col)].width = 10

        ws.cell(row=current_row, column=money_col, value="金額").alignment = align_center
        ws.cell(row=current_row, column=sign_col, value="學生簽名").alignment = align_center
        ws.cell(row=current_row, column=check_col, value="已收錢").alignment = align_center

        for c in range(1, check_col + 1):
            ws.cell(row=current_row, column=c).border = border
        current_row += 1

        headers = ["班級", "座號", "姓名", "學號"]
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=current_row, column=i, value=h)
            c.alignment = align_center
            c.border = border

        for p in products:
            map_info = product_col_map[p.id]
            start_c = map_info['start']
            for s_idx, size in enumerate(map_info['sizes']):
                c = ws.cell(row=current_row, column=start_c + s_idx, value=size)
                c.alignment = align_center
                c.border = border
                c.fill = fill_grey if map_info['color_idx'] % 2 == 0 else fill_white

        for c in range(money_col, check_col + 1):
            ws.cell(row=current_row, column=c).border = border
        current_row += 1

        for order in class_orders_list:
            ws.row_dimensions[current_row].height = 25

            user = order.user
            ws.cell(row=current_row, column=1, value=user.class_name).alignment = align_center
            ws.cell(row=current_row, column=2, value=user.seat_number).alignment = align_center
            ws.cell(row=current_row, column=3, value=user.student_name).alignment = align_center
            ws.cell(row=current_row, column=4, value=user.login_id).alignment = align_center

            student_total = 0

            for p in products:
                map_info = product_col_map[p.id]
                start_c = map_info['start']
                for s_idx in range(len(map_info['sizes'])):
                    c = ws.cell(row=current_row, column=start_c + s_idx, value="")
                    c.border = border
                    c.fill = fill_grey if map_info['color_idx'] % 2 == 0 else fill_white

            for item in order.items:
                map_info = product_col_map.get(item.product_id)
                if map_info:
                    try:
                        s_idx = map_info['sizes'].index(item.item_size)
                        c = ws.cell(row=current_row, column=map_info['start'] + s_idx, value=item.quantity)
                        c.alignment = align_center
                    except ValueError:
                        pass
                    student_total += item.product.price * item.quantity

            class_total_money += student_total

            c_money = ws.cell(row=current_row, column=money_col, value=student_total)
            c_money.alignment = align_center
            c_money.border = border

            ws.cell(row=current_row, column=sign_col, value="").border = border
            ws.cell(row=current_row, column=check_col, value="").border = border

            for c in range(1, 5):
                ws.cell(row=current_row, column=c).border = border

            current_row += 1

        ws.cell(row=current_row, column=1, value="(接著是備用空格)").font = Font(color="7688A7")
        current_row += 1

        for _ in range(3):
            ws.row_dimensions[current_row].height = 25
            for c in range(1, check_col + 1):
                ws.cell(row=current_row, column=c).border = border
            current_row += 1

        current_row += 1
        ws.merge_cells(start_row=current_row, start_column=money_col - 3, end_row=current_row, end_column=check_col)
        total_cell = ws.cell(row=current_row, column=money_col - 3, value=f"最終金額：{class_total_money} + （     ） = （     ）")
        total_cell.font = Font(bold=True)
        total_cell.alignment = Alignment(horizontal='right', vertical='center')

        medium_side = Side(style='medium')
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=4, max_col=4):
            for cell in row:
                cell.border = Border(left=cell.border.left, top=cell.border.top, bottom=cell.border.bottom, right=medium_side)
        for row in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=medium_side)

        ws.oddFooter.right.text = "&14part1：各班分頁報表(班代用) - 第&P頁"
        ws.evenFooter.right.text = "&14part1：各班分頁報表(班代用) - 第&P頁"

    if default_sheet.title == "Sheet":
        wb.remove(default_sheet)

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name=f'part1：各班分頁報表(班代用){get_excel_suffix()}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# 匯出 excel-2 全校總表
@app.route('/manager/order/export_excel_all')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_export_excel_all():
    products = OrderProduct.query.all()

    orders = db.session.query(OrderMain).join(OrderUser)\
             .filter(OrderMain.status == '已送出')\
             .order_by(OrderUser.class_name, OrderUser.seat_number).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "全校總表"

    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    fill_grey = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    ws.freeze_panes = 'E3'

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 6
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 12

    headers = ["班級", "座號", "姓名", "學號"]
    for i, h in enumerate(headers, 1):
        ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
        cell = ws.cell(row=1, column=i, value=h)
        cell.alignment = align_center
        cell.border = border
        ws.cell(row=2, column=i).border = border

    col_idx = 5
    product_col_map = {}
    size_totals = {}

    for idx, p in enumerate(products):
        sizes = p.sizes.split(',')

        ws.merge_cells(start_row=1, start_column=col_idx, end_row=1, end_column=col_idx + len(sizes) - 1)
        p_cell = ws.cell(row=1, column=col_idx, value=f"{p.name} ({p.price}元)")
        p_cell.alignment = align_center
        p_cell.border = border
        p_cell.font = Font(bold=True)

        product_col_map[p.id] = {'start': col_idx, 'sizes': sizes, 'color_idx': idx}

        for s_idx, size in enumerate(sizes):
            c_col = col_idx + s_idx
            s_cell = ws.cell(row=2, column=c_col, value=size)
            s_cell.alignment = align_center
            s_cell.border = border
            s_cell.fill = fill_grey if idx % 2 == 0 else fill_white
            ws.column_dimensions[openpyxl.utils.get_column_letter(c_col)].width = 5
            size_totals[(p.id, size)] = 0

        col_idx += len(sizes)

    money_col = col_idx
    ws.column_dimensions[openpyxl.utils.get_column_letter(money_col)].width = 14
    ws.merge_cells(start_row=1, start_column=money_col, end_row=2, end_column=money_col)
    m_cell = ws.cell(row=1, column=money_col, value="個人總金額")
    m_cell.alignment = align_center
    m_cell.border = border
    m_cell.font = Font(bold=True)
    ws.cell(row=2, column=money_col).border = border

    current_row = 3
    grand_total_money = 0

    for order in orders:
        user = order.user
        ws.cell(row=current_row, column=1, value=user.class_name).alignment = align_center
        ws.cell(row=current_row, column=2, value=user.seat_number).alignment = align_center
        ws.cell(row=current_row, column=3, value=user.student_name).alignment = align_center
        ws.cell(row=current_row, column=4, value=user.login_id).alignment = align_center

        for c in range(1, 5):
            ws.cell(row=current_row, column=c).border = border

        student_total = 0

        for p in products:
            map_info = product_col_map[p.id]
            start_c = map_info['start']
            for s_idx in range(len(map_info['sizes'])):
                c = ws.cell(row=current_row, column=start_c + s_idx, value="")
                c.border = border
                c.fill = fill_grey if map_info['color_idx'] % 2 == 0 else fill_white

        for item in order.items:
            map_info = product_col_map.get(item.product_id)
            if map_info:
                try:
                    s_idx = map_info['sizes'].index(item.item_size)
                    c = ws.cell(row=current_row, column=map_info['start'] + s_idx, value=item.quantity)
                    c.alignment = align_center
                    size_totals[(item.product_id, item.item_size)] += item.quantity
                except ValueError:
                    pass
                student_total += item.product.price * item.quantity

        c_money = ws.cell(row=current_row, column=money_col, value=student_total)
        c_money.alignment = align_center
        c_money.border = border

        grand_total_money += student_total
        current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    c_total_label = ws.cell(row=current_row, column=1, value="全校總計")
    c_total_label.alignment = align_center
    c_total_label.font = Font(bold=True, size=12)
    for c in range(1, 5):
        ws.cell(row=current_row, column=c).border = border

    for p in products:
        map_info = product_col_map[p.id]
        start_c = map_info['start']
        for s_idx, size in enumerate(map_info['sizes']):
            total_qty = size_totals[(p.id, size)]
            c = ws.cell(row=current_row, column=start_c + s_idx, value=total_qty)
            c.alignment = align_center
            c.border = border
            c.font = Font(bold=True, color="E53E3E")
            c.fill = fill_grey if map_info['color_idx'] % 2 == 0 else fill_white

    c_grand_total = ws.cell(row=current_row, column=money_col, value=grand_total_money)
    c_grand_total.alignment = align_center
    c_grand_total.border = border
    c_grand_total.font = Font(bold=True, color="E53E3E", size=12)

    # ================= 加上分隔粗線 (不加頁尾) =================
    medium_side = Side(style='medium')
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=4, max_col=4):
        for cell in row:
            cell.border = Border(left=cell.border.left, top=cell.border.top, bottom=cell.border.bottom, right=medium_side)
    for row in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=medium_side)
    # ==========================================================

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name=f'other：全校總表{get_excel_suffix()}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# 匯出 excel-3 各班對帳簽收總表 (單一工作表，高一/高二/高三分頁列印)
@app.route('/manager/order/export_excel_class_collection')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_export_excel_class_collection():
    orders = db.session.query(OrderMain).join(OrderUser)\
             .filter(OrderMain.status == '已送出')\
             .order_by(OrderUser.class_name).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "各班收費簽收表"

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = '1:1'

    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')

    headers = ["班級", "原始金額", "+ (            )", " = 最終金額", "實收", "是否收完？", "班代簽名", "班聯負責人簽名"]
    col_widths = [10, 15, 18, 18, 12, 12, 22, 22]

    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.alignment = align_center
        cell.border = border
        cell.font = Font(bold=True, size=12)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = col_widths[i-1]

    ws.row_dimensions[1].height = 30

    current_row = 2
    previous_grade = None
    orders_by_class = groupby(orders, key=lambda o: o.user.class_name)

    for class_name, class_orders in orders_by_class:
        current_grade = str(class_name)[0] if class_name else ""
        if previous_grade is not None and current_grade != previous_grade:
            ws.row_breaks.append(Break(id=current_row - 1))
        previous_grade = current_grade

        class_total_money = 0
        for order in class_orders:
            for item in order.items:
                class_total_money += item.product.price * item.quantity

        c_class = ws.cell(row=current_row, column=1, value=class_name)
        c_money = ws.cell(row=current_row, column=2, value=class_total_money)
        c_plus = ws.cell(row=current_row, column=3, value="+ （        ）")
        c_final = ws.cell(row=current_row, column=4, value=" = （        ）")
        c_actual = ws.cell(row=current_row, column=5, value="")
        c_status = ws.cell(row=current_row, column=6, value="")
        c_sign1 = ws.cell(row=current_row, column=7, value="")
        c_sign2 = ws.cell(row=current_row, column=8, value="")

        for c in [c_class, c_money, c_plus, c_final, c_actual, c_status, c_sign1, c_sign2]:
            c.alignment = align_center
            c.border = border
            c.font = Font(size=12)

        c_class.font = Font(bold=True, size=12)
        ws.row_dimensions[current_row].height = 40
        current_row += 1

    # ================= 加上分隔粗線與頁尾設定 =================
    medium_side = Side(style='medium')
    for row in ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=medium_side)

    ws.oddFooter.right.text = "&14part1：各班收費報表(班聯用) - 第&P頁"
    ws.evenFooter.right.text = "&14part1：各班收費報表(班聯用) - 第&P頁"
    # ==========================================================

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name=f'part1：各班收費報表(班聯用){get_excel_suffix()}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# 匯出 excel-4 每班 Size 總額與發放簽收表 (一項產品一頁)
@app.route('/manager/order/export_excel_class_size')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_export_excel_class_size():
    products = OrderProduct.query.all()

    orders = db.session.query(OrderMain).join(OrderUser)\
             .filter(OrderMain.status == '已送出').all()

    agg_data = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))

    for order in orders:
        user = order.user
        for item in order.items:
            agg_data[user.class_name][item.product_id][item.item_size] += item.quantity

    all_classes = sorted(list(agg_data.keys()))
    categories = {"高一": [], "高二": [], "高三": [], "國中_處室": []}
    for c in all_classes:
        if c.startswith('1'): categories["高一"].append(c)
        elif c.startswith('2'): categories["高二"].append(c)
        elif c.startswith('3'): categories["高三"].append(c)
        else: categories["國中_處室"].append(c)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "發放簽收表"

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.scale = 100
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.freeze_panes = 'B3'
    ws.print_title_rows = '1:2'
    ws.print_title_cols = 'A:A'

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    fill_grey = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20
    ws.column_dimensions['A'].width = 12

    c_class1 = ws.cell(row=1, column=1, value="班級 / 處室")
    c_class2 = ws.cell(row=2, column=1, value="")
    c_class1.alignment = align_center
    c_class1.border = thin_border
    c_class1.font = Font(bold=True)
    ws.merge_cells("A1:A2")
    c_class2.border = thin_border

    col_idx = 2
    product_col_map = {}

    for idx, p in enumerate(products):
        sizes = p.sizes.split(',')
        start_col = col_idx

        ws.merge_cells(start_row=1, start_column=col_idx, end_row=1, end_column=col_idx + len(sizes) - 1)
        p_cell = ws.cell(row=1, column=col_idx, value=f"{p.name} ({p.price}元)")
        p_cell.alignment = align_center
        p_cell.border = thin_border
        p_cell.font = Font(bold=True, size=12)

        for s_idx, size in enumerate(sizes):
            c_col = col_idx + s_idx
            s_cell = ws.cell(row=2, column=c_col, value=size)
            s_cell.alignment = align_center
            s_cell.border = thin_border
            s_cell.fill = fill_grey if idx % 2 == 0 else fill_white
            ws.column_dimensions[openpyxl.utils.get_column_letter(c_col)].width = 6

        col_idx += len(sizes)

        sign1_col = col_idx
        sign2_col = col_idx + 1

        ws.merge_cells(start_row=1, start_column=sign1_col, end_row=2, end_column=sign1_col)
        cell_s1 = ws.cell(row=1, column=sign1_col, value="班代簽名")
        cell_s1.alignment = align_center
        cell_s1.border = thin_border
        cell_s1.font = Font(bold=True)
        ws.cell(row=2, column=sign1_col).border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(sign1_col)].width = 16

        ws.merge_cells(start_row=1, start_column=sign2_col, end_row=2, end_column=sign2_col)
        cell_s2 = ws.cell(row=1, column=sign2_col, value="班聯負責人簽名")
        cell_s2.alignment = align_center
        cell_s2.border = thin_border
        cell_s2.font = Font(bold=True)
        ws.cell(row=2, column=sign2_col).border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(sign2_col)].width = 18

        product_col_map[p.id] = {'start': start_col, 'sizes': sizes, 'color_idx': idx, 'sign_cols': (sign1_col, sign2_col)}

        if idx != len(products) - 1:
            ws.col_breaks.append(Break(id=sign2_col))

        col_idx += 2

    current_row = 3
    ordered_classes = categories["高一"] + categories["高二"] + categories["高三"] + categories["國中_處室"]

    for class_name in ordered_classes:
        c_label = ws.cell(row=current_row, column=1, value=class_name)
        c_label.alignment = align_center
        c_label.border = thin_border
        c_label.font = Font(bold=True)

        for p in products:
            map_info = product_col_map[p.id]
            start_c = map_info['start']

            for s_idx, size in enumerate(map_info['sizes']):
                qty = agg_data[class_name][p.id][size]
                c = ws.cell(row=current_row, column=start_c + s_idx, value=qty if qty > 0 else "")
                c.alignment = align_center
                c.border = thin_border
                c.fill = fill_grey if map_info['color_idx'] % 2 == 0 else fill_white

            ws.cell(row=current_row, column=map_info['sign_cols'][0], value="").border = thin_border
            ws.cell(row=current_row, column=map_info['sign_cols'][1], value="").border = thin_border

        current_row += 1

    ws_confirm = wb.create_sheet(title="全校完成確認表")
    ws_confirm.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws_confirm.page_margins.left = 0.5
    ws_confirm.page_margins.top = 0.5

    conf_row = 2
    title_cell = ws_confirm.cell(row=conf_row, column=1, value="✅ 全校班級 / 處室 發放完成總確認表")
    title_cell.font = Font(bold=True, size=16, color="003A7C")
    conf_row += 2

    for cat_name in ["高一", "高二", "高三", "國中_處室"]:
        c_list = categories[cat_name]
        if not c_list: continue

        cat_cell = ws_confirm.cell(row=conf_row, column=1, value=f"【 {cat_name.replace('_', '/')} 】")
        cat_cell.font = Font(bold=True, size=14)
        conf_row += 1

        for i, c_name in enumerate(c_list):
            cell = ws_confirm.cell(row=conf_row, column=i + 1, value=c_name)
            cell.alignment = align_center
            cell.border = thin_border
            cell.fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')

        conf_row += 1

        ws_confirm.row_dimensions[conf_row].height = 25
        for i in range(len(c_list)):
            cell = ws_confirm.cell(row=conf_row, column=i + 1, value="")
            cell.border = thin_border

        conf_row += 2

    # ================= 加上分隔粗線與頁尾設定 =================
    medium_side = Side(style='medium')
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=1):
        for cell in row:
            cell.border = Border(left=cell.border.left, top=cell.border.top, bottom=cell.border.bottom, right=medium_side)
    for row in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=medium_side)

    ws.oddFooter.right.text = "&14part3：各班size總額(班聯用) - 第&P頁"
    ws.evenFooter.right.text = "&14part3：各班size總額(班聯用) - 第&P頁"

    ws_confirm.oddFooter.right.text = "&14part3：各班size總額(班聯用)_確認表 - 第&P頁"
    ws_confirm.evenFooter.right.text = "&14part3：各班size總額(班聯用)_確認表 - 第&P頁"
    # ==========================================================

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name=f'part3：各班size總額(班聯用){get_excel_suffix()}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# 匯出 excel-5 個人發放衣服對應表 (各班分頁，每項商品最右側有已發送打勾)
@app.route('/manager/order/export_excel_student_delivery')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_export_excel_student_delivery():
    products = OrderProduct.query.all()

    orders = db.session.query(OrderMain).join(OrderUser)\
             .filter(OrderMain.status == '已送出')\
             .order_by(OrderUser.class_name, OrderUser.seat_number).all()

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    fill_grey = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_check = PatternFill(start_color='FFFBEB', end_color='FFFBEB', fill_type='solid')

    orders_by_class = groupby(orders, key=lambda o: o.user.class_name)

    for class_name, class_orders in orders_by_class:
        class_orders_list = list(class_orders)

        ws = wb.create_sheet(title=f"{class_name}班")

        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 85

        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

        ws.print_title_rows = '1:2'
        ws.print_title_cols = 'A:D'
        ws.freeze_panes = 'E3'

        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 6
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 12

        headers = ["班級", "座號", "姓名", "學號"]
        for i, h in enumerate(headers, 1):
            ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
            cell = ws.cell(row=1, column=i, value=h)
            cell.alignment = align_center
            cell.border = thin_border
            ws.cell(row=2, column=i).border = thin_border

        col_idx = 5
        product_col_map = {}

        for idx, p in enumerate(products):
            sizes = p.sizes.split(',')

            ws.merge_cells(start_row=1, start_column=col_idx, end_row=1, end_column=col_idx + len(sizes))

            p_cell = ws.cell(row=1, column=col_idx, value=f"{p.name}")
            p_cell.alignment = align_center
            p_cell.border = thin_border
            p_cell.font = Font(bold=True)

            for s_idx, size in enumerate(sizes):
                c_col = col_idx + s_idx
                s_cell = ws.cell(row=2, column=c_col, value=size)
                s_cell.alignment = align_center
                s_cell.border = thin_border
                s_cell.fill = fill_grey if idx % 2 == 0 else fill_white
                ws.column_dimensions[openpyxl.utils.get_column_letter(c_col)].width = 5

            check_col = col_idx + len(sizes)
            chk_cell = ws.cell(row=2, column=check_col, value="已發送")
            chk_cell.alignment = align_center
            chk_cell.border = thin_border
            chk_cell.font = Font(bold=True, color="D97706")
            chk_cell.fill = fill_check
            ws.column_dimensions[openpyxl.utils.get_column_letter(check_col)].width = 8

            product_col_map[p.id] = {'start': col_idx, 'sizes': sizes, 'check_col': check_col, 'color_idx': idx}
            col_idx += len(sizes) + 1

        current_row = 3

        for order in class_orders_list:
            ws.row_dimensions[current_row].height = 25

            user = order.user
            ws.cell(row=current_row, column=1, value=user.class_name).alignment = align_center
            ws.cell(row=current_row, column=2, value=user.seat_number).alignment = align_center
            ws.cell(row=current_row, column=3, value=user.student_name).alignment = align_center
            ws.cell(row=current_row, column=4, value=user.login_id).alignment = align_center

            for c in range(1, 5):
                ws.cell(row=current_row, column=c).border = thin_border

            for p in products:
                map_info = product_col_map[p.id]
                start_c = map_info['start']

                for s_idx in range(len(map_info['sizes'])):
                    c = ws.cell(row=current_row, column=start_c + s_idx, value="")
                    c.border = thin_border
                    c.fill = fill_grey if map_info['color_idx'] % 2 == 0 else fill_white

                chk_c = ws.cell(row=current_row, column=map_info['check_col'], value="")
                chk_c.border = thin_border
                chk_c.fill = fill_check

            for item in order.items:
                map_info = product_col_map.get(item.product_id)
                if map_info:
                    try:
                        s_idx = map_info['sizes'].index(item.item_size)
                        c = ws.cell(row=current_row, column=map_info['start'] + s_idx, value=item.quantity)
                        c.alignment = align_center
                        c.font = Font(bold=True)
                    except ValueError:
                        pass

            current_row += 1

        # ================= 加上分隔粗線與頁尾設定 =================
        medium_side = Side(style='medium')
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=4, max_col=4):
            for cell in row:
                cell.border = Border(left=cell.border.left, top=cell.border.top, bottom=cell.border.bottom, right=medium_side)
        for row in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=medium_side)

        ws.oddFooter.right.text = "&14part3：個人發放衣服對應表(班代用) - 第&P頁"
        ws.evenFooter.right.text = "&14part3：個人發放衣服對應表(班代用) - 第&P頁"
        # ==========================================================

    if default_sheet.title == "Sheet":
        wb.remove(default_sheet)

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name=f'part3：個人發放衣服對應表(班代用){get_excel_suffix()}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# 匯出 excel-6 廠商訂購總表 (縱軸產品、橫軸Size)
@app.route('/manager/order/export_excel_manufacturer')
@manager_student_required
@manager_dept_required(['manager'])
def manager_order_export_excel_manufacturer():
    products = OrderProduct.query.all()

    # 抓取「已送出」訂單
    orders = db.session.query(OrderMain).filter(OrderMain.status == '已送出').all()

    # 1. 蒐集所有出現過的 Size (並保持上架時的順序)
    all_sizes = []
    for p in products:
        for s in p.sizes.split(','):
            if s not in all_sizes:
                all_sizes.append(s)

    # 2. 彙整資料：agg_data[product_id][size] = 總數量
    agg_data = defaultdict(lambda: defaultdict(int))
    for order in orders:
        for item in order.items:
            agg_data[item.product_id][item.item_size] += item.quantity

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "廠商訂購單"

    # 基礎樣式
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    fill_header = PatternFill(start_color='CBD5E0', end_color='CBD5E0', fill_type='solid')
    fill_light = PatternFill(start_color='F7FAFC', end_color='F7FAFC', fill_type='solid')

    ws.column_dimensions['A'].width = 25
    ws.row_dimensions[1].height = 25

    # ========== 建立橫軸 (X軸) 表頭 ==========
    c_corner = ws.cell(row=1, column=1, value="產品名稱 \\ Size")
    c_corner.alignment = align_center
    c_corner.border = thin_border
    c_corner.fill = fill_header
    c_corner.font = Font(bold=True)

    # 印出所有 Size
    for col_idx, size in enumerate(all_sizes, start=2):
        c_size = ws.cell(row=1, column=col_idx, value=size)
        c_size.alignment = align_center
        c_size.border = thin_border
        c_size.fill = fill_header
        c_size.font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 8

    # 最右側：單品總計欄
    total_col_idx = len(all_sizes) + 2
    c_row_total_head = ws.cell(row=1, column=total_col_idx, value="單品總計")
    c_row_total_head.alignment = align_center
    c_row_total_head.border = thin_border
    c_row_total_head.fill = PatternFill(start_color='FEFCBF', end_color='FEFCBF', fill_type='solid')
    c_row_total_head.font = Font(bold=True, color="B7791F")
    ws.column_dimensions[openpyxl.utils.get_column_letter(total_col_idx)].width = 12

    # ========== 填寫縱軸 (Y軸) 與資料 ==========
    current_row = 2
    for p in products:
        # 產品名稱
        c_name = ws.cell(row=current_row, column=1, value=f"{p.name}")
        c_name.alignment = align_center
        c_name.border = thin_border
        c_name.font = Font(bold=True)

        row_total = 0

        # 填入各尺寸數量
        for col_idx, size in enumerate(all_sizes, start=2):
            qty = agg_data[p.id][size]
            c_qty = ws.cell(row=current_row, column=col_idx, value=qty if qty > 0 else "")
            c_qty.alignment = align_center
            c_qty.border = thin_border
            if qty > 0:
                c_qty.font = Font(bold=True)

            row_total += qty

        # 填入單品總計
        c_row_total = ws.cell(row=current_row, column=total_col_idx, value=row_total)
        c_row_total.alignment = align_center
        c_row_total.border = thin_border
        c_row_total.font = Font(bold=True, color="B7791F")
        c_row_total.fill = PatternFill(start_color='FEFCBF', end_color='FEFCBF', fill_type='solid')

        current_row += 1

    # ========== 最下方：尺寸總計列 ==========
    c_col_total_head = ws.cell(row=current_row, column=1, value="【所有款式】各尺寸總計")
    c_col_total_head.alignment = align_center
    c_col_total_head.border = thin_border
    c_col_total_head.fill = fill_light
    c_col_total_head.font = Font(bold=True)

    grand_total = 0
    for col_idx, size in enumerate(all_sizes, start=2):
        # 加總該尺寸在所有產品中的數量
        col_total = sum(agg_data[p.id][size] for p in products)
        c_col_total = ws.cell(row=current_row, column=col_idx, value=col_total)
        c_col_total.alignment = align_center
        c_col_total.border = thin_border
        c_col_total.fill = fill_light
        c_col_total.font = Font(bold=True, color="2B6CB0")

        grand_total += col_total

    # 右下角：全站總數量
    c_grand_total = ws.cell(row=current_row, column=total_col_idx, value=grand_total)
    c_grand_total.alignment = align_center
    c_grand_total.border = thin_border
    c_grand_total.fill = PatternFill(start_color='FEEBC8', end_color='FEEBC8', fill_type='solid')
    c_grand_total.font = Font(bold=True, color="C05621", size=14)

    import io
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name=f'part2：廠商訂購總表.xlsx{get_excel_suffix()}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# ==========================================
# 第 12 區：啟動伺服器
# ==========================================
if __name__ == '__main__':
    app.run()
